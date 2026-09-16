import io
import logging
import re
import uuid
from datetime import datetime, timezone
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Response, Query, status
from sqlalchemy.orm import Session
import openpyxl

from app.api.deps import get_current_user, require_data_source_access
from app.core.database import get_db
from app.models.user import User
from app.models.reference_data import (
    InternationalMarketBrand,
    IqviaExtract,
    RegisteredNotInUse,
    WhoInnRegistry,
)
from app.repositories.screening import ScreeningRepository
from app.repositories.settings import SettingsRepository
from app.schemas.reference_data import (
    DataSourceListResponse,
    DataSourceStatus,
    DataSourceToggleRequest,
    IqviaBulkDeleteRequest,
    IqviaPaginatedResponse,
    IqviaRecordCreate,
    IqviaRecordResponse,
    IqviaRecordUpdate,
    ReferenceDataStatus,
    UploadResponse,
    WhoInnPaginatedResponse,
    WhoInnRecordCreate,
    WhoInnRecordUpdate,
    WhoInnRecordResponse,
    InternationalMarketPaginatedResponse,
    InternationalMarketRecordCreate,
    InternationalMarketRecordUpdate,
    InternationalMarketRecordResponse,
    RegisteredNotInUsePaginatedResponse,
    RegisteredNotInUseRecordCreate,
    RegisteredNotInUseRecordUpdate,
    RegisteredNotInUseRecordResponse,
)
from app.services.market_check import google_search_configured
from app.services.tabular_import import (
    normalize_name,
    parse_international_market_xlsx,
    parse_iqvia_file,
    parse_registered_not_in_use_xlsx,
)
from app.services.who_inn_import import normalize_inn_name, parse_who_inn_pdf

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/reference-data", tags=["Reference Data"])

# Upload endpoints in this file only stream a workbook/PDF fully into memory
# (no streaming parse), so an unbounded upload is a memory-exhaustion vector.
MAX_UPLOAD_BYTES = 50 * 1024 * 1024

_SPREADSHEET_CONTENT_TYPES = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel.sheet.binary.macroenabled.12",
    "application/octet-stream",
)


def _require_xlsx(file: UploadFile) -> None:
    if file.content_type not in _SPREADSHEET_CONTENT_TYPES and not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="File must be an .xlsx workbook")


def _require_spreadsheet(file: UploadFile) -> None:
    fn = (file.filename or "").lower()
    if not (fn.endswith(".xlsx") or fn.endswith(".xlsb") or fn.endswith(".xlsm")):
        if file.content_type not in _SPREADSHEET_CONTENT_TYPES:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="File must be an .xlsx or .xlsb workbook")


@router.get("/status", response_model=ReferenceDataStatus)
def get_reference_data_status(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    repo = ScreeningRepository(db)
    return ReferenceDataStatus(
        who_inn_row_count=repo.who_inn_row_count(),
        iqvia_row_count=repo.iqvia_row_count(),
        registered_not_in_use_row_count=repo.registered_not_in_use_row_count(),
        international_market_row_count=repo.international_market_row_count(),
    )


# ---------------------------------------------------------------------------
# Toggleable data sources — Connect / Disconnect
# ---------------------------------------------------------------------------

def _build_data_source_statuses(repo: ScreeningRepository, toggles: dict) -> list[DataSourceStatus]:
    who_count = repo.who_inn_row_count()
    iqvia_active_count = repo.iqvia_row_count(active_only=True)
    iqvia_total_count = repo.iqvia_row_count(active_only=False)
    who_enabled = toggles.get("who_inn_enabled", True)
    iqvia_enabled = toggles.get("iqvia_enabled", True)
    epharmacy_enabled = toggles.get("epharmacy_enabled", True)
    google_enabled = toggles.get("google_search_enabled", True)
    google_configured = google_search_configured()

    return [
        DataSourceStatus(
            id="who_inn",
            name="WHO INN Registry",
            category="Regulatory Databases",
            description="International Nonproprietary Names. A name identical to a registered INN is an automatic knockout.",
            enabled=who_enabled,
            connected=who_enabled,
            detail=(
                f"{who_count} entries loaded from upload." if who_count > 0
                else "No entries loaded — upload the WHO INN list to enable this check."
            ) if who_enabled else "Disabled. WHO INN is not checked during screening.",
        ),
        DataSourceStatus(
            id="iqvia",
            name="IQVIA Extract",
            category="Market Intelligence",
            description="Licensed IQVIA market-data extract covering market brands, molecules, sales, and volumes.",
            enabled=iqvia_enabled,
            connected=iqvia_enabled and iqvia_active_count > 0,
            detail=(
                f"Connected ({iqvia_active_count} active records loaded)." if iqvia_active_count > 0
                else f"{iqvia_total_count} records loaded (all inactive)." if iqvia_total_count > 0
                else "No data loaded."
            ) if iqvia_enabled else "Disabled. IQVIA is not checked during screening.",
        ),
        DataSourceStatus(
            id="epharmacy",
            name="E-Pharmacy Platforms",
            category="Market Intelligence",
            description="Live scrape of 1mg, PharmEasy, Apollo Pharmacy & Netmeds for active listings under the candidate name.",
            enabled=epharmacy_enabled,
            connected=epharmacy_enabled,
            detail=(
                "Live scrape active: 1mg, PharmEasy, Apollo Pharmacy, Netmeds." if epharmacy_enabled
                else "Disabled. E-pharmacy platforms are not checked during screening."
            ),
        ),
        DataSourceStatus(
            id="google_search",
            name="Google Search",
            category="Market Intelligence",
            description="General web presence check via Google Custom Search: market/regulatory mentions outside e-pharmacy listings.",
            enabled=google_enabled,
            connected=google_enabled and google_configured,
            detail=(
                ("Configured." if google_configured else "Not configured. Set GOOGLE_API_KEY/GOOGLE_CSE_ID.")
                if google_enabled else "Disabled. Google Search is not checked during screening."
            ),
        ),
    ]


@router.get("/data-sources", response_model=DataSourceListResponse)
def list_data_sources(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    repo = ScreeningRepository(db)
    toggles = SettingsRepository(db).get_data_source_toggles()
    return DataSourceListResponse(sources=_build_data_source_statuses(repo, toggles))


@router.put("/data-sources/{source_id}", response_model=DataSourceStatus)
def toggle_data_source(
    source_id: str,
    payload: DataSourceToggleRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    key_map = {
        "who_inn": "who_inn_enabled",
        "iqvia": "iqvia_enabled",
        "epharmacy": "epharmacy_enabled",
        "google_search": "google_search_enabled",
    }
    if source_id not in key_map:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Unknown data source '{source_id}'")

    settings_repo = SettingsRepository(db)
    toggles = settings_repo.set_data_source_toggle(key_map[source_id], payload.enabled)
    logger.info("[DATA SOURCE TOGGLE] %s set to enabled=%s by %s", source_id, payload.enabled, current_user.email)

    repo = ScreeningRepository(db)
    statuses = _build_data_source_statuses(repo, toggles)
    return next(s for s in statuses if s.id == source_id)


# ---------------------------------------------------------------------------
# WHO INN — PDF upload
# ---------------------------------------------------------------------------

@router.post("/who-inn/upload", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_who_inn_list(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):

    if file.content_type not in ("application/pdf", "application/octet-stream") and not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="File must be a PDF")

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File exceeds maximum allowed size (50MB)")

    try:
        parsed = parse_who_inn_pdf(file_bytes)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception:
        logger.exception("Unexpected failure parsing WHO INN PDF upload %r", file.filename)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read this PDF. Is it the expected WHO INN list template?")

    rows = [
        {"inn_name": r["inn_name"], "normalized_name": normalize_inn_name(r["inn_name"]),
         "who_publication_reference": r["who_publication_reference"]}
        for r in parsed
    ]

    repo = ScreeningRepository(db)
    imported = repo.replace_all_who_inn(rows)
    logger.info("[WHO INN IMPORT] Replaced who_inn_registry with %d rows from %r (by %s)",
                imported, file.filename, current_user.email)

    return UploadResponse(
        rows_imported=imported,
        message=f"Imported {imported} WHO INN entries from '{file.filename}'.",
    )


@router.post("/who-inn/upload-append", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_who_inn_list_append(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    """Adds entries from the uploaded PDF to the existing WHO INN table,
    skipping any name that's already present, instead of replacing it."""

    if file.content_type not in ("application/pdf", "application/octet-stream") and not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="File must be a PDF")

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File exceeds maximum allowed size (50MB)")

    try:
        parsed = parse_who_inn_pdf(file_bytes)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception:
        logger.exception("Unexpected failure parsing WHO INN PDF upload %r", file.filename)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read this PDF. Is it the expected WHO INN list template?")

    rows = [
        {"inn_name": r["inn_name"], "normalized_name": normalize_inn_name(r["inn_name"]),
         "who_publication_reference": r["who_publication_reference"]}
        for r in parsed
    ]

    repo = ScreeningRepository(db)
    added, skipped = repo.append_who_inn(rows)
    logger.info("[WHO INN IMPORT] Appended %d rows (%d skipped as duplicates) from %r (by %s)",
                added, skipped, file.filename, current_user.email)

    return UploadResponse(
        rows_imported=added,
        rows_skipped=skipped,
        message=f"Added {added} new WHO INN entries from '{file.filename}'"
                + (f" ({skipped} duplicate name(s) skipped)." if skipped else "."),
    )


# ---------------------------------------------------------------------------
# MASTER DATA — WHO INN Registry (CRUD + Template + Export)
# ---------------------------------------------------------------------------

@router.get("/who-inn/records", response_model=WhoInnPaginatedResponse)
def get_who_inn_records(
    q: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    repo = ScreeningRepository(db)
    items, total = repo.get_who_inn_records_paginated(q=q, page=page, page_size=page_size)
    total_pages = (total + page_size - 1) // page_size if total > 0 else 1
    return WhoInnPaginatedResponse(items=items, total=total, page=page, page_size=page_size, total_pages=total_pages)


@router.post("/who-inn/records", response_model=WhoInnRecordResponse, status_code=status.HTTP_201_CREATED)
def create_who_inn_record(
    payload: WhoInnRecordCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    if not payload.inn_name or not payload.inn_name.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="INN name is required.")

    record = WhoInnRegistry(
        inn_name=payload.inn_name.strip(),
        normalized_name=normalize_inn_name(payload.inn_name),
        who_publication_reference=payload.who_publication_reference.strip() if payload.who_publication_reference else None,
        chembl_id=payload.chembl_id.strip() if payload.chembl_id else None,
        molecule_type=payload.molecule_type.strip() if payload.molecule_type else None,
        created_at=datetime.now(timezone.utc),
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


@router.put("/who-inn/records/{record_id}", response_model=WhoInnRecordResponse)
def update_who_inn_record(
    record_id: uuid.UUID,
    payload: WhoInnRecordUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    record = db.query(WhoInnRegistry).filter(WhoInnRegistry.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Record not found.")

    if payload.inn_name is not None:
        if not payload.inn_name.strip():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="INN name cannot be empty.")
        record.inn_name = payload.inn_name.strip()
        record.normalized_name = normalize_inn_name(payload.inn_name)
    if payload.who_publication_reference is not None:
        record.who_publication_reference = payload.who_publication_reference.strip() if payload.who_publication_reference else None
    if payload.chembl_id is not None:
        record.chembl_id = payload.chembl_id.strip() if payload.chembl_id else None
    if payload.molecule_type is not None:
        record.molecule_type = payload.molecule_type.strip() if payload.molecule_type else None

    db.commit()
    db.refresh(record)
    return record


@router.delete("/who-inn/records/{record_id}", status_code=status.HTTP_200_OK)
def delete_who_inn_record(
    record_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    record = db.query(WhoInnRegistry).filter(WhoInnRegistry.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Record not found.")
    db.delete(record)
    db.commit()
    return {"message": "Record deleted successfully."}


@router.get("/who-inn/template")
def download_who_inn_template(
    current_user: User = Depends(get_current_user),
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WHO INN Template"
    headers = ["Sr No.", "INN Name", "W.H.O Publication Reference"]
    ws.append(headers)

    ws.append([1, "ABACAVIR", "List 77 (1997)"])
    ws.append([2, "METFORMIN", "List 4 (1956)"])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=WHO_INN_Template.xlsx"},
    )


@router.get("/who-inn/export")
def export_who_inn_master_data(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    records = db.query(WhoInnRegistry).order_by(WhoInnRegistry.inn_name.asc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "WHO INN Registry"
    headers = ["Sr No.", "INN Name", "W.H.O Publication Reference", "Created Date"]
    ws.append(headers)

    for idx, r in enumerate(records, start=1):
        ws.append([
            idx,
            r.inn_name,
            r.who_publication_reference or "",
            r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=WHO_INN_Registry_{datetime.now().strftime('%Y%m%d')}.xlsx"},
    )


# ---------------------------------------------------------------------------
# MASTER DATA — International Market Brands (CRUD + Template + Export + Bulk)
# ---------------------------------------------------------------------------

@router.get("/international-market/records", response_model=InternationalMarketPaginatedResponse)
def get_international_market_records(
    q: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    repo = ScreeningRepository(db)
    items, total = repo.get_international_market_records_paginated(q=q, page=page, page_size=page_size)
    total_pages = (total + page_size - 1) // page_size if total > 0 else 1
    return InternationalMarketPaginatedResponse(items=items, total=total, page=page, page_size=page_size, total_pages=total_pages)


@router.post("/international-market/records", response_model=InternationalMarketRecordResponse, status_code=status.HTTP_201_CREATED)
def create_international_market_record(
    payload: InternationalMarketRecordCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    if not payload.brand_name or not payload.brand_name.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Brand name (Mark) is required.")

    record = InternationalMarketBrand(
        brand_name=payload.brand_name.strip(),
        normalized_name=normalize_name(payload.brand_name),
        active_ingredient=payload.active_ingredient.strip() if payload.active_ingredient else None,
        country=payload.country.strip() if payload.country else None,
        as_of_date=payload.as_of_date,
        created_at=datetime.now(timezone.utc),
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


@router.put("/international-market/records/{record_id}", response_model=InternationalMarketRecordResponse)
def update_international_market_record(
    record_id: uuid.UUID,
    payload: InternationalMarketRecordUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    record = db.query(InternationalMarketBrand).filter(InternationalMarketBrand.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Record not found.")

    if payload.brand_name is not None:
        if not payload.brand_name.strip():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Brand name cannot be empty.")
        record.brand_name = payload.brand_name.strip()
        record.normalized_name = normalize_name(payload.brand_name)
    if payload.active_ingredient is not None:
        record.active_ingredient = payload.active_ingredient.strip() if payload.active_ingredient else None
    if payload.country is not None:
        record.country = payload.country.strip() if payload.country else None
    if payload.as_of_date is not None:
        record.as_of_date = payload.as_of_date

    db.commit()
    db.refresh(record)
    return record


@router.delete("/international-market/records/{record_id}", status_code=status.HTTP_200_OK)
def delete_international_market_record(
    record_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    record = db.query(InternationalMarketBrand).filter(InternationalMarketBrand.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Record not found.")
    db.delete(record)
    db.commit()
    return {"message": "Record deleted successfully."}


@router.get("/international-market/template")
def download_international_market_template(
    current_user: User = Depends(get_current_user),
):
    """Generates the exact template matching info/data_sourrces_uploads/international.xlsx:
    Headers: Sr No., Mark, Molecule"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Sr No.", "Mark", "Molecule"])

    # Sample rows matching the official template
    ws.append([1, "TICAPLET", "TICAGRELOR"])
    ws.append([2, "TICAFAST", "TICAGRELOR"])
    ws.append([3, "ANGRELOR", "TICAGRELOR"])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=International_Markets_Template.xlsx"},
    )


@router.get("/international-market/export")
def export_international_market_master_data(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    records = db.query(InternationalMarketBrand).order_by(InternationalMarketBrand.brand_name.asc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Sr No.", "Mark", "Molecule"])

    for idx, r in enumerate(records, start=1):
        ws.append([
            idx,
            r.brand_name,
            r.active_ingredient or "",
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=International_Market_Brands_{datetime.now().strftime('%Y%m%d')}.xlsx"},
    )


@router.post("/international-market/upload", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_international_market(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    _require_xlsx(file)

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File exceeds maximum allowed size (50MB)")

    try:
        parsed = parse_international_market_xlsx(file_bytes)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception:
        logger.exception("Unexpected failure parsing International Market upload %r", file.filename)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read this workbook. Is it the expected template?")

    repo = ScreeningRepository(db)
    imported = repo.replace_all_international_market(parsed)
    logger.info("[INTERNATIONAL-MARKET IMPORT] Replaced table with %d rows from %r (by %s)",
                imported, file.filename, current_user.email)

    return UploadResponse(
        rows_imported=imported,
        message=f"Imported {imported} International Market Brand entries from '{file.filename}'.",
    )


@router.post("/international-market/upload-append", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_international_market_append(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    """Adds entries to the existing table, skipping any name that's already
    present, instead of replacing it."""
    _require_xlsx(file)

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File exceeds maximum allowed size (50MB)")

    try:
        parsed = parse_international_market_xlsx(file_bytes)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception:
        logger.exception("Unexpected failure parsing International Market upload %r", file.filename)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read this workbook. Is it the expected template?")

    repo = ScreeningRepository(db)
    added, skipped = repo.append_international_market(parsed)
    logger.info("[INTERNATIONAL-MARKET IMPORT] Appended %d rows (%d skipped as duplicates) from %r (by %s)",
                added, skipped, file.filename, current_user.email)

    return UploadResponse(
        rows_imported=added,
        rows_skipped=skipped,
        message=f"Added {added} new International Market Brand entries from '{file.filename}'"
                + (f" ({skipped} duplicate name(s) skipped)." if skipped else "."),
    )


# ---------------------------------------------------------------------------
# MASTER DATA — Registered but Not in Use (CRUD + Template + Export + Bulk)
# ---------------------------------------------------------------------------

@router.get("/registered-not-in-use/records", response_model=RegisteredNotInUsePaginatedResponse)
def get_registered_not_in_use_records(
    q: Optional[str] = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    repo = ScreeningRepository(db)
    items, total = repo.get_registered_not_in_use_records_paginated(q=q, page=page, page_size=page_size)
    total_pages = (total + page_size - 1) // page_size if total > 0 else 1
    return RegisteredNotInUsePaginatedResponse(items=items, total=total, page=page, page_size=page_size, total_pages=total_pages)


@router.post("/registered-not-in-use/records", response_model=RegisteredNotInUseRecordResponse, status_code=status.HTTP_201_CREATED)
def create_registered_not_in_use_record(
    payload: RegisteredNotInUseRecordCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    if not payload.brand_name or not payload.brand_name.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Brand Name (TradeMark Name) is required.")

    record = RegisteredNotInUse(
        brand_name=payload.brand_name.strip(),
        normalized_name=normalize_name(payload.brand_name),
        trademark_class=payload.trademark_class or 5,
        application_number=payload.application_number.strip() if payload.application_number else None,
        application_date=payload.application_date,
        status=payload.status.strip() if payload.status else "Registered",
        valid_till=payload.valid_till,
        remarks=payload.remarks.strip() if payload.remarks else None,
        as_of_date=payload.as_of_date,
        created_at=datetime.now(timezone.utc),
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


@router.put("/registered-not-in-use/records/{record_id}", response_model=RegisteredNotInUseRecordResponse)
def update_registered_not_in_use_record(
    record_id: uuid.UUID,
    payload: RegisteredNotInUseRecordUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    record = db.query(RegisteredNotInUse).filter(RegisteredNotInUse.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Record not found.")

    if payload.brand_name is not None:
        if not payload.brand_name.strip():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Brand name cannot be empty.")
        record.brand_name = payload.brand_name.strip()
        record.normalized_name = normalize_name(payload.brand_name)
    if payload.trademark_class is not None:
        record.trademark_class = payload.trademark_class
    if payload.application_number is not None:
        record.application_number = payload.application_number.strip() if payload.application_number else None
    if payload.application_date is not None:
        record.application_date = payload.application_date
    if payload.status is not None:
        record.status = payload.status.strip() if payload.status else None
    if payload.valid_till is not None:
        record.valid_till = payload.valid_till
    if payload.remarks is not None:
        record.remarks = payload.remarks.strip() if payload.remarks else None
    if payload.as_of_date is not None:
        record.as_of_date = payload.as_of_date

    db.commit()
    db.refresh(record)
    return record


@router.delete("/registered-not-in-use/records/{record_id}", status_code=status.HTTP_200_OK)
def delete_registered_not_in_use_record(
    record_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    record = db.query(RegisteredNotInUse).filter(RegisteredNotInUse.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Record not found.")
    db.delete(record)
    db.commit()
    return {"message": "Record deleted successfully."}


@router.get("/registered-not-in-use/template")
def download_registered_not_in_use_template(
    current_user: User = Depends(get_current_user),
):
    """Generates the exact template matching info/data_sourrces_uploads/register_not_used.xlsx:
    Headers: Sl No, TradeMark Name, Class, Appl No, Appl Date, TMR STATUS, Valid till, Description, E-commerce URL"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = ["Sl No", "TradeMark Name", "Class", "Appl No", "Appl Date", "TMR STATUS", "Valid till", "Description", "E-commerce URL"]
    ws.append(headers)

    # Sample rows matching the official template
    ws.append([1, "A2CLEAR", 5, "5693827", "2022-11-22", "Registered", "2032-11-22", "N-1 Recommended to maintain this mark", ""])
    ws.append([2, "ACTRIL", 5, "564924", "1992-01-03", "Registered", "2026-01-03", "Dormant registered mark", ""])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Registered_Not_In_Use_Template.xlsx"},
    )


@router.get("/registered-not-in-use/export")
def export_registered_not_in_use_master_data(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    records = db.query(RegisteredNotInUse).order_by(RegisteredNotInUse.brand_name.asc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = ["Sl No", "TradeMark Name", "Class", "Appl No", "Appl Date", "TMR STATUS", "Valid till", "Description", "E-commerce URL"]
    ws.append(headers)

    for idx, r in enumerate(records, start=1):
        ws.append([
            idx,
            r.brand_name,
            r.trademark_class or 5,
            r.application_number or "",
            r.application_date.strftime("%Y-%m-%d") if r.application_date else "",
            r.status or "Registered",
            r.valid_till.strftime("%Y-%m-%d") if r.valid_till else "",
            r.remarks or "",
            "",
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Registered_Not_In_Use_{datetime.now().strftime('%Y%m%d')}.xlsx"},
    )


@router.post("/registered-not-in-use/upload", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_registered_not_in_use(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    _require_xlsx(file)

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File exceeds maximum allowed size (50MB)")

    try:
        parsed = parse_registered_not_in_use_xlsx(file_bytes)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception:
        logger.exception("Unexpected failure parsing Registered-Not-in-Use upload %r", file.filename)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read this workbook. Is it the expected template?")

    repo = ScreeningRepository(db)
    imported = repo.replace_all_registered_not_in_use(parsed)
    logger.info("[REGISTERED-NOT-IN-USE IMPORT] Replaced table with %d rows from %r (by %s)",
                imported, file.filename, current_user.email)

    return UploadResponse(
        rows_imported=imported,
        message=f"Imported {imported} Registered-but-Not-in-Use entries from '{file.filename}'.",
    )


@router.post("/registered-not-in-use/upload-append", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_registered_not_in_use_append(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    """Adds entries to the existing table, skipping any name that's already
    present, instead of replacing it."""
    _require_xlsx(file)

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File exceeds maximum allowed size (50MB)")

    try:
        parsed = parse_registered_not_in_use_xlsx(file_bytes)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception:
        logger.exception("Unexpected failure parsing Registered-Not-in-Use upload %r", file.filename)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read this workbook. Is it the expected template?")

    repo = ScreeningRepository(db)
    added, skipped = repo.append_registered_not_in_use(parsed)
    logger.info("[REGISTERED-NOT-IN-USE IMPORT] Appended %d rows (%d skipped as duplicates) from %r (by %s)",
                added, skipped, file.filename, current_user.email)

    return UploadResponse(
        rows_imported=added,
        rows_skipped=skipped,
        message=f"Added {added} new Registered-but-Not-in-Use entries from '{file.filename}'"
                + (f" ({skipped} duplicate name(s) skipped)." if skipped else "."),
    )


# ---------------------------------------------------------------------------
# MASTER DATA — IQVIA Extract (CRUD + Upload + Template + Export + Bulk)
# ---------------------------------------------------------------------------

@router.get("/iqvia/records", response_model=IqviaPaginatedResponse)
def get_iqvia_records(
    q: Optional[str] = Query(None),
    is_active: Optional[bool] = Query(None),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    repo = ScreeningRepository(db)
    items, total = repo.get_iqvia_records_paginated(q=q, is_active=is_active, page=page, page_size=page_size)
    total_pages = (total + page_size - 1) // page_size if total > 0 else 1
    return IqviaPaginatedResponse(
        items=items,
        total=total,
        page=page,
        page_size=page_size,
        total_pages=total_pages,
    )


@router.post("/iqvia/records", response_model=IqviaRecordResponse, status_code=status.HTTP_201_CREATED)
def create_iqvia_record(
    payload: IqviaRecordCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    if not payload.brand_ims or not payload.brand_ims.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="BRAND_IMS (Brand Name) is required.")

    val_gr = payload.val_gr_pct
    if val_gr is None and payload.val_mat_current is not None and payload.val_mat_prev is not None and payload.val_mat_prev != 0:
        val_gr = round(((payload.val_mat_current - payload.val_mat_prev) / payload.val_mat_prev) * 100, 2)

    un_gr = payload.un_gr_pct
    if un_gr is None and payload.un_mat_current is not None and payload.un_mat_prev is not None and payload.un_mat_prev != 0:
        un_gr = round(((payload.un_mat_current - payload.un_mat_prev) / payload.un_mat_prev) * 100, 2)

    norm_name = re.sub(r"[^a-z0-9]", "", payload.brand_ims.strip().lower())
    record = IqviaExtract(
        brand_ims=payload.brand_ims.strip(),
        normalized_name=norm_name,
        molecules=payload.molecules.strip() if payload.molecules else None,
        atc_iv=payload.atc_iv.strip() if payload.atc_iv else None,
        company=payload.company.strip() if payload.company else None,
        product_launch=payload.product_launch.strip() if payload.product_launch else None,
        val_mat_current=payload.val_mat_current,
        val_mat_prev=payload.val_mat_prev,
        val_gr_pct=val_gr,
        un_mat_current=payload.un_mat_current,
        un_mat_prev=payload.un_mat_prev,
        un_gr_pct=un_gr,
        no_of_mol=payload.no_of_mol,
        plain_comb=payload.plain_comb.strip() if payload.plain_comb else None,
        is_active=payload.is_active if payload.is_active is not None else True,
        as_of_date=payload.as_of_date,
        created_at=datetime.now(timezone.utc),
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return record


@router.put("/iqvia/records/{record_id}", response_model=IqviaRecordResponse)
def update_iqvia_record(
    record_id: uuid.UUID,
    payload: IqviaRecordUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    record = db.query(IqviaExtract).filter(IqviaExtract.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Record not found.")

    if payload.brand_ims is not None:
        if not payload.brand_ims.strip():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Brand IMS cannot be empty.")
        record.brand_ims = payload.brand_ims.strip()
        record.normalized_name = re.sub(r"[^a-z0-9]", "", payload.brand_ims.strip().lower())
    if payload.molecules is not None:
        record.molecules = payload.molecules.strip() if payload.molecules else None
    if payload.atc_iv is not None:
        record.atc_iv = payload.atc_iv.strip() if payload.atc_iv else None
    if payload.company is not None:
        record.company = payload.company.strip() if payload.company else None
    if payload.product_launch is not None:
        record.product_launch = payload.product_launch.strip() if payload.product_launch else None
    if payload.val_mat_current is not None:
        record.val_mat_current = payload.val_mat_current
    if payload.val_mat_prev is not None:
        record.val_mat_prev = payload.val_mat_prev
    if payload.val_gr_pct is not None:
        record.val_gr_pct = payload.val_gr_pct
    elif record.val_mat_current is not None and record.val_mat_prev is not None and record.val_mat_prev != 0:
        record.val_gr_pct = round(((record.val_mat_current - record.val_mat_prev) / record.val_mat_prev) * 100, 2)
    if payload.un_mat_current is not None:
        record.un_mat_current = payload.un_mat_current
    if payload.un_mat_prev is not None:
        record.un_mat_prev = payload.un_mat_prev
    if payload.un_gr_pct is not None:
        record.un_gr_pct = payload.un_gr_pct
    elif record.un_mat_current is not None and record.un_mat_prev is not None and record.un_mat_prev != 0:
        record.un_gr_pct = round(((record.un_mat_current - record.un_mat_prev) / record.un_mat_prev) * 100, 2)
    if payload.no_of_mol is not None:
        record.no_of_mol = payload.no_of_mol
    if payload.plain_comb is not None:
        record.plain_comb = payload.plain_comb.strip() if payload.plain_comb else None
    if payload.is_active is not None:
        record.is_active = payload.is_active
    if payload.as_of_date is not None:
        record.as_of_date = payload.as_of_date

    db.commit()
    db.refresh(record)
    return record


@router.delete("/iqvia/records/{record_id}", status_code=status.HTTP_200_OK)
def delete_iqvia_record(
    record_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    record = db.query(IqviaExtract).filter(IqviaExtract.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Record not found.")
    db.delete(record)
    db.commit()
    return {"message": "Record deleted successfully."}


@router.post("/iqvia/records/bulk-delete", status_code=status.HTTP_200_OK)
def bulk_delete_iqvia_records(
    payload: IqviaBulkDeleteRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    repo = ScreeningRepository(db)
    deleted = repo.bulk_delete_iqvia(ids=payload.ids, clear_all=payload.clear_all)
    return {"message": f"Deleted {deleted} IQVIA record(s).", "deleted_count": deleted}


@router.put("/iqvia/records/{record_id}/toggle-active", response_model=IqviaRecordResponse)
def toggle_iqvia_record_active(
    record_id: uuid.UUID,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    record = db.query(IqviaExtract).filter(IqviaExtract.id == record_id).first()
    if not record:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Record not found.")
    record.is_active = not record.is_active
    db.commit()
    db.refresh(record)
    return record


@router.get("/iqvia/template")
def download_iqvia_template(
    current_user: User = Depends(get_current_user),
):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IQVIA_Extract"
    headers = [
        "Brand IMS", "Molecules", "ATC IV", "Company", "Product Launch",
        "Val MAT Jun 26", "Val MAT Jun 25", "Val GR Pct",
        "UN MAT Jun 26", "UN MAT Jun 25", "UN GR Pct",
        "No Of Mol", "Plain Comb"
    ]
    ws.append(headers)
    ws.append([
        "AUGMENTIN", "AMOXICILLIN+CLAVULANIC ACID", "J01C2", "GSK", "199805",
        450000000.0, 410000000.0, 9.76, 25000000.0, 23500000.0, 6.38, 2, "COMB"
    ])
    ws.append([
        "VOLINI", "DICLOFENAC+LINSEED OIL+MENTHOL+METHYL SALICYLATE", "M02AA", "SUN PHARMA", "199408",
        320000000.0, 290000000.0, 10.34, 18000000.0, 16800000.0, 7.14, 4, "COMB"
    ])
    ws.append([
        "LIPITOR", "ATORVASTATIN", "C10AA", "PFIZER", "199701",
        150000000.0, 140000000.0, 7.14, 12000000.0, 11500000.0, 4.35, 1, "PLAIN"
    ])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=IQVIA_Master_Data_Template.xlsx"},
    )


@router.get("/iqvia/export")
def export_iqvia_master_data(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    records = db.query(IqviaExtract).order_by(IqviaExtract.brand_ims.asc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IQVIA_Master_Data"
    headers = [
        "Brand IMS", "Molecules", "ATC IV", "Company", "Product Launch",
        "Val MAT Jun 26", "Val MAT Jun 25", "Val GR Pct",
        "UN MAT Jun 26", "UN MAT Jun 25", "UN GR Pct",
        "No Of Mol", "Plain Comb", "Status", "Created At"
    ]
    ws.append(headers)
    for r in records:
        ws.append([
            r.brand_ims,
            r.molecules or "",
            r.atc_iv or "",
            r.company or "",
            r.product_launch or "",
            r.val_mat_current if r.val_mat_current is not None else "",
            r.val_mat_prev if r.val_mat_prev is not None else "",
            r.val_gr_pct if r.val_gr_pct is not None else "",
            r.un_mat_current if r.un_mat_current is not None else "",
            r.un_mat_prev if r.un_mat_prev is not None else "",
            r.un_gr_pct if r.un_gr_pct is not None else "",
            r.no_of_mol if r.no_of_mol is not None else "",
            r.plain_comb or "",
            "Active" if r.is_active else "Inactive",
            r.created_at.strftime("%Y-%m-%d %H:%M") if r.created_at else "",
        ])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=IQVIA_Master_Data_{datetime.now().strftime('%Y%m%d')}.xlsx"},
    )


@router.post("/iqvia/upload", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_iqvia(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    _require_spreadsheet(file)

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File exceeds maximum allowed size (50MB)")

    try:
        parsed = parse_iqvia_file(file_bytes, filename=file.filename or "")
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception:
        logger.exception("Unexpected failure parsing IQVIA upload %r", file.filename)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read this workbook. Is it the expected template?")

    repo = ScreeningRepository(db)
    imported = repo.replace_all_iqvia(parsed)
    logger.info("[IQVIA IMPORT] Replaced table with %d rows from %r (by %s)",
                imported, file.filename, current_user.email)

    return UploadResponse(
        rows_imported=imported,
        message=f"Imported {imported} IQVIA entries from '{file.filename}'.",
    )


@router.post("/iqvia/upload-append", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_iqvia_append(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_data_source_access),
):
    """Adds entries to the existing table, skipping any brand name that's
    already present, instead of replacing it."""
    _require_spreadsheet(file)

    file_bytes = await file.read()
    if not file_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded file is empty")
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File exceeds maximum allowed size (50MB)")

    try:
        parsed = parse_iqvia_file(file_bytes, filename=file.filename or "")
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception:
        logger.exception("Unexpected failure parsing IQVIA upload %r", file.filename)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not read this workbook. Is it the expected template?")

    repo = ScreeningRepository(db)
    added, skipped = repo.append_iqvia(parsed)
    logger.info("[IQVIA IMPORT] Appended %d rows (%d skipped as duplicates) from %r (by %s)",
                added, skipped, file.filename, current_user.email)

    return UploadResponse(
        rows_imported=added,
        rows_skipped=skipped,
        message=f"Added {added} new IQVIA entries from '{file.filename}'"
                + (f" ({skipped} duplicate name(s) skipped)." if skipped else "."),
    )
